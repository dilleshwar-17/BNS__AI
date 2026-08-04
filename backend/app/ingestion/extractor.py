"""Text extraction for supported source file types (Phase 4)."""

from __future__ import annotations

from pathlib import Path

from app.core.exceptions import DocumentProcessingFailedError


def extract_text(path: str | Path) -> str:
    """Extract raw text from a PDF, DOCX, or TXT file."""
    p = Path(path)
    ext = p.suffix.lower()
    try:
        if ext == ".pdf":
            return _extract_pdf(p)
        if ext in {".docx", ".doc"}:
            return _extract_docx(p)
        if ext == ".txt":
            return p.read_text(encoding="utf-8", errors="replace")
        raise DocumentProcessingFailedError(f"Unsupported file type for extraction: {ext}")
    except DocumentProcessingFailedError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise DocumentProcessingFailedError(f"Text extraction failed: {exc}") from exc


def _extract_pdf(p: Path) -> str:
    try:
        import fitz  # PyMuPDF
    except ImportError as exc:  # pragma: no cover
        raise DocumentProcessingFailedError("PyMuPDF is required for PDF extraction.") from exc

    pages: list[str] = []
    with fitz.open(p) as doc:
        for page in doc:
            pages.append(page.get_text("text"))
    return "\n".join(pages)


def _extract_docx(p: Path) -> str:
    try:
        from docx import Document  # python-docx
    except ImportError as exc:  # pragma: no cover
        raise DocumentProcessingFailedError("python-docx is required for DOCX extraction.") from exc
    d = Document(str(p))
    parts: list[str] = []
    for para in d.paragraphs:
        if para.text:
            parts.append(para.text)
    for table in d.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts)
