import re
from io import BytesIO

from app.core.logging import get_logger

logger = get_logger(__name__)


class DraftExporter:
    async def export_txt(self, content: str, title: str) -> bytes:
        header = f"{title}\n{'=' * len(title)}\n\n"
        return (header + content).encode("utf-8")

    async def export_markdown(self, content: str, title: str) -> bytes:
        md = f"# {title}\n\n{content}\n"
        return md.encode("utf-8")

    async def export_docx(self, content: str, title: str) -> bytes:
        try:
            from docx import Document
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.shared import Inches, Pt

            buffer = BytesIO()
            doc = Document()

            style = doc.styles["Normal"]
            font = style.font
            font.name = "Times New Roman"
            font.size = Pt(12)

            section = doc.sections[0]
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1.25)
            section.right_margin = Inches(1.25)

            title_para = doc.add_heading(title, level=1)
            title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

            for line in content.split("\n"):
                stripped = line.strip()
                if not stripped:
                    continue
                if stripped.startswith("**") and stripped.endswith("**"):
                    p = doc.add_paragraph()
                    run = p.add_run(stripped.strip("**"))
                    run.bold = True
                    run.font.size = Pt(12)
                elif stripped.startswith("# "):
                    doc.add_heading(stripped[2:], level=1)
                elif stripped.startswith("## "):
                    doc.add_heading(stripped[3:], level=2)
                elif stripped.startswith("### "):
                    doc.add_heading(stripped[4:], level=3)
                elif re.match(r"^\d+[.)]", stripped):
                    p = doc.add_paragraph(stripped)
                    p.paragraph_format.space_after = Pt(2)
                elif stripped.startswith("* ") or stripped.startswith("- "):
                    p = doc.add_paragraph(stripped, style="List Bullet")
                else:
                    p = doc.add_paragraph(stripped)
                    p.paragraph_format.space_after = Pt(4)

            doc.save(buffer)
            buffer.seek(0)
            return buffer.getvalue()
        except ImportError:
            logger.warning("docx_export_failed_import")
            return await self.export_txt(content, title)
        except Exception as e:
            logger.error("docx_export_error", error=str(e))
            return await self.export_txt(content, title)

    async def export_pdf(self, content: str, title: str) -> bytes:
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
            from reportlab.lib.units import inch
            from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

            buffer = BytesIO()
            doc = SimpleDocTemplate(
                buffer, pagesize=A4, topMargin=inch, bottomMargin=inch, leftMargin=inch * 1.25, rightMargin=inch * 1.25
            )

            styles = getSampleStyleSheet()
            title_style = ParagraphStyle("Title2", parent=styles["Title"], spaceAfter=12)
            normal_style = ParagraphStyle(
                "CustomNormal", parent=styles["Normal"], fontName="Times-Roman", fontSize=12, leading=15
            )

            elements = []
            elements.append(Paragraph(title, title_style))
            elements.append(Spacer(1, 12))

            for line in content.split("\n"):
                stripped = line.strip()
                if not stripped:
                    elements.append(Spacer(1, 6))
                    continue
                if stripped.startswith("**") and stripped.endswith("**"):
                    bold_style = ParagraphStyle("Bold", parent=normal_style, fontName="Times-Bold")
                    elements.append(Paragraph(stripped.strip("**"), bold_style))
                elif stripped.startswith("# ") or stripped.startswith("## ") or stripped.startswith("### "):
                    heading = re.sub(r"^#+\s*", "", stripped)
                    hs = ParagraphStyle(
                        "Heading", parent=normal_style, fontName="Times-Bold", fontSize=14, spaceAfter=8, spaceBefore=12
                    )
                    elements.append(Paragraph(heading, hs))
                else:
                    elements.append(Paragraph(stripped, normal_style))

            doc.build(elements)
            buffer.seek(0)
            return buffer.getvalue()
        except ImportError:
            logger.warning("pdf_export_failed_import")
            return await self.export_txt(content, title)
        except Exception as e:
            logger.error("pdf_export_error", error=str(e))
            return await self.export_txt(content, title)


draft_exporter = DraftExporter()
