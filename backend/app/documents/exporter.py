from io import BytesIO


class DocumentExporter:
    async def to_txt(self, content: str) -> bytes:
        return content.encode("utf-8")

    async def to_docx(self, content: str, title: str) -> bytes:
        try:
            from docx import Document
            from docx.shared import Pt

            buffer = BytesIO()
            doc = Document()
            style = doc.styles["Normal"]
            style.font.name = "Times New Roman"
            style.font.size = Pt(12)
            doc.add_heading(title, level=1)
            for line in content.split("\n"):
                if line.strip():
                    doc.add_paragraph(line.strip())
            doc.save(buffer)
            buffer.seek(0)
            return buffer.getvalue()
        except ImportError:
            return content.encode("utf-8")
